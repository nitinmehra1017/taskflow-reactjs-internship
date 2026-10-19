"""Rebuild the Week 1 report using python-docx from the bundled document runtime."""

from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "ReactJS_Internship_Week1_Project_Plan.docx"
doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.27)
section.page_height = Inches(11.69)
section.top_margin = Inches(0.78)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)
section.header_distance = Inches(0.3)
section.footer_distance = Inches(0.3)
section.different_first_page_header_footer = True

for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"]:
    style = doc.styles[name]
    style.font.name = "Arial"
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_after = Pt(8)

normal = doc.styles["Normal"]
normal.font.size = Pt(10.5)
normal.paragraph_format.line_spacing = 1.15
normal.paragraph_format.widow_control = True
for name, size in [("Title", 29), ("Heading 1", 21), ("Heading 2", 13), ("Heading 3", 11)]:
    doc.styles[name].font.size = Pt(size)
    doc.styles[name].font.bold = name != "Title"
    doc.styles[name].paragraph_format.keep_with_next = True
doc.styles["Subtitle"].font.size = Pt(13)
doc.styles["Subtitle"].font.italic = False

header = section.header.paragraphs[0]
header.text = "TaskFlow    /    YuvaIntern ReactJS Internship    /    Week 1"
header.runs[0].font.size = Pt(8)
header.runs[0].font.color.rgb = RGBColor(0, 0, 0)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run("Project setup and planning    |    ").font.size = Pt(8)
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)

doc.core_properties.title = "TaskFlow ReactJS Internship Week 1 Project Plan"
doc.core_properties.subject = "YuvaIntern Junior ReactJS Developer Internship"
doc.core_properties.author = "TaskFlow project workspace"
doc.core_properties.keywords = "ReactJS, Webpack, Week 1, YuvaIntern, TaskFlow"


def p(text, style=None):
    return doc.add_paragraph(text, style=style)


def h(text, level=2):
    return doc.add_heading(text, level=level)


def page(title):
    doc.add_page_break()
    h(title, 1)


def bullet(text):
    par = p("•  " + text)
    par.paragraph_format.left_indent = Inches(0.1)
    par.paragraph_format.first_line_indent = Inches(-0.1)
    return par


def code(text, size=8.7):
    for line in text.strip("\n").splitlines():
        par = p(line or " ")
        par.paragraph_format.space_after = Pt(0)
        par.paragraph_format.line_spacing = 1.02
        par.paragraph_format.left_indent = Inches(0.1)
        par.paragraph_format.right_indent = Inches(0.08)
        par.paragraph_format.keep_with_next = False
        shade = OxmlElement("w:shd")
        shade.set(qn("w:fill"), "F2F4F5")
        par._p.get_or_add_pPr().append(shade)
        for run in par.runs:
            run.font.name = "Courier New"
            run.font.size = Pt(size)
    p("").paragraph_format.space_after = Pt(2)


def table(headers, rows, widths, size=9):
    tab = doc.add_table(rows=1, cols=len(headers))
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    tab.autofit = False
    for col, width in zip(tab.columns, widths):
        col.width = Inches(width)
    for cell, title, width in zip(tab.rows[0].cells, headers, widths):
        cell.width = Inches(width)
        cell.text = title
    repeat = OxmlElement("w:tblHeader")
    tab.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for cell, val, width in zip(tab.add_row().cells, row, widths):
            cell.width = Inches(width)
            cell.text = str(val)
    for ri, row in enumerate(tab.rows):
        cant_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant_split)
        for ci, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for side in ["top", "left", "bottom", "right"]:
                el = OxmlElement("w:" + side)
                el.set(qn("w:w"), "85" if side in ["top", "bottom"] else "95")
                el.set(qn("w:type"), "dxa")
                margins.append(el)
            props.append(margins)
            borders = OxmlElement("w:tcBorders")
            for side in ["top", "left", "bottom", "right"]:
                el = OxmlElement("w:" + side)
                el.set(qn("w:val"), "single")
                el.set(qn("w:sz"), "4")
                el.set(qn("w:color"), "D9D9D9")
                borders.append(el)
            props.append(borders)
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "243D48" if ri == 0 else ("F0F4F6" if ri % 2 == 0 else "FFFFFF"))
            props.append(shade)
            for par in cell.paragraphs:
                par.paragraph_format.space_after = Pt(0)
                par.paragraph_format.line_spacing = 1.08
                if widths[ci] < 0.8:
                    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in par.runs:
                    run.font.size = Pt(size)
                    run.font.bold = ri == 0
                    run.font.color.rgb = RGBColor(255, 255, 255) if ri == 0 else RGBColor(0, 0, 0)
    p("").paragraph_format.space_after = Pt(1)
    return tab


# 1 — Title and scope statement.
p("YuvaIntern Junior ReactJS Developer Internship", "Subtitle")
p("").paragraph_format.space_after = Pt(35)
p("TaskFlow ReactJS Internship Week 1 Project Plan", "Title")
p("Project setup and planning", "Subtitle")
p("Week 1 of a six week internship")
p("Prepared on 19 October 2026")
p("").paragraph_format.space_after = Pt(26)
h("Purpose of this submission")
p("TaskFlow now has a functional React starter, a custom Webpack build pipeline, a shared application shell, and a documented plan for the next five weeks. This report explains the delivered foundation, the decisions behind it, and the verification evidence available for internship review.")
p("The Week 1 boundary is intentional. Task CRUD, persistence, filters, and automated feature tests are planned for later weeks. Dashboard figures describe the project structure; they do not represent user activity.")
h("Submission contents")
bullet("Source project with an npm lockfile, development scripts, routing, organized styles, and scalable folders.")
bullet("Architecture, component contracts, coding standards, six week roadmap, and a planned 35 hour work breakdown.")
bullet("Genuine command logs and this Word report, with explicit limits on browser verification, screenshots, and GitHub publication.")
p("Project directory: taskflow-reactjs-internship")
p("Report file: ReactJS_Internship_Week1_Project_Plan.docx")

# 2 — Introduction and scope.
page("Project introduction and objectives")
p("TaskFlow is a personal task management dashboard intended to help a user capture work, organize priorities, and track completion. Its six week progression begins with a reliable development foundation so later features can be added without mixing routing, presentation, state, and storage concerns.")
h("Week 1 objectives")
bullet("Initialize a real JavaScript and JSX React project with working start, build, lint, format, and format check commands.")
bullet("Configure a transparent Webpack 5 and Babel toolchain, including HTML generation, CSS loading, asset modules, source maps, and a development server.")
bullet("Deliver an attractive responsive shell with a header, navigation, overview, task preview, setup page, and unknown route recovery.")
bullet("Establish component boundaries, naming rules, quality practices, and an accessible starting point.")
bullet("Document future scope, record actual command results, and prepare a reviewable internship submission.")
h("Proposed features and scope")
table(["Capability", "Week 1 status", "Later delivery"], [
    ["Shared shell and routing", "Implemented", "Extend as features grow"],
    ["Overview dashboard", "Structural overview and roadmap", "Derived task statistics in Week 4"],
    ["Task creation and editing", "Planned empty state only", "Week 2"],
    ["Complete and remove tasks", "Not implemented", "Week 2"],
    ["Shared state and storage", "Folder boundaries documented", "Week 3"],
    ["Search, filters and preferences", "Not implemented", "Week 4"],
    ["Automated feature tests", "Strategy documented", "Week 5"],
    ["Production hosting", "Root deployment prepared", "Week 6"],
], [2.0, 2.3, 2.4])
p("Authentication, collaboration, remote APIs, notifications, and paid services are outside the current internship plan. They should be evaluated separately if requirements change.")

# 3 — Stack.
page("Technology stack and justification")
p("The implementation uses maintained Node.js release lines and a custom toolchain suited to learning the build process. Exact versions are recorded in package-lock.json and evidence/environment.log. The versions below are the resolved versions actually installed for this submission.")
table(["Technology", "Resolved version", "Reason for selection"], [
    ["Node.js and npm", "24.19.0 / 11.21.1", "Modern runtime and reproducible dependency installation"],
    ["React and React DOM", "19.3.1", "Component composition, declarative UI, and JSX"],
    ["React Router DOM", "7.18.4", "Nested routes, shared layouts, and active navigation"],
    ["Webpack", "5.112.0", "Explicit modules, assets, bundling, and production optimization"],
    ["Webpack CLI / dev server", "6.0.1 / 5.2.7", "Build commands, local server, HMR, and route fallback"],
    ["Babel core and presets", "7.29.7", "Environment transforms and automatic React JSX runtime"],
    ["ESLint / React X", "10.13.0 / 5.28.2", "Maintained flat configuration and React checks"],
    ["React Hooks plugin", "7.1.1", "Official hook rules and dependency checks"],
    ["Prettier", "3.9.10", "Consistent formatting independent of lint rules"],
    ["HTMLWebpackPlugin", "5.6.8", "Generate the entry HTML with hashed bundle references"],
    ["CSS loader / style loader", "7.1.5 / 4.0.0", "Resolve CSS imports and inject styles during development"],
    ["MiniCssExtractPlugin", "2.10.2", "Separate production CSS for browser caching"],
], [1.9, 1.4, 3.4], size=8.8)
h("Compatibility decisions")
p("Node.js 24 is selected by .nvmrc; package engines also permit maintained Node.js 22.14 and newer within the 22 release line. Only Node 24.19.0 was executed in this environment. Babel 7 is retained as a compatible toolchain line with babel-loader 10; the installed package metadata reported no Babel deprecation.")
p("The first install warned that ESLint 9 was unsupported. It was replaced with ESLint 10, React X, and the official Hooks plugin. The final clean installation and lint command succeeded. Accessibility is reviewed through source practices and the planned manual/browser strategy; an accessibility lint plugin is not claimed.")

# 4 — Directory and architecture.
page("Development approach and architecture")
p("The application uses small functional components. App owns route declarations; AppLayout owns the shared shell; pages compose reusable UI. Future task state will live behind a provider only when multiple routes need it. Services will own persistence, while utilities remain pure. This avoids introducing unused code before the task model is implemented.")
h("Project directory structure")
code("""taskflow-reactjs-internship/
  public/index.html
  src/
    index.jsx                  React root and BrowserRouter
    App.jsx                    Nested routes
    components/                Header Sidebar Icon StatCard PageHeading
    layouts/AppLayout.jsx      Shared shell and Outlet
    pages/                     Dashboard Tasks Settings NotFound
    hooks/useRouteFocus.js     Title and navigation focus
    context/README.md          Future cross route state
    services/README.md         Future storage boundary
    utils/README.md            Future pure helpers
    styles/                    global.css shell.css dashboard.css
    assets/                    Original SVG mark and asset notes
  docs/                        Architecture standards roadmap report
  evidence/                    Actual logs and evidence notes
  package.json                 Scripts and dependency ranges
  package-lock.json            Exact dependency resolution
  webpack.config.cjs           Development and production build
  babel.config.cjs             Environment and JSX compilation
  eslint.config.mjs            Flat lint configuration
  .prettierrc.json              Formatting rules
  .editorconfig                Editor defaults
  .gitignore                   Secrets and generated file exclusions
  .nvmrc                       Node.js 24
  submission-description.txt  Internship submission field text""", 8.7)
h("Routing and ownership")
p("The root route renders the Overview dashboard. /tasks renders the planned task empty state, /settings renders development guidance, and the wildcard route provides a recovery link. BrowserRouter wraps App, and a shared parent route places the selected page inside Outlet.")
p("NavLink supplies active route semantics. useRouteFocus updates the browser title and focuses main after a pathname change. Organized plain CSS separates design tokens, shell layout, and dashboard/page styles. CSS Modules can be introduced later if ownership becomes difficult; no CSS framework is required.")

# 5 — Hierarchy and contracts.
page("Component hierarchy and responsibilities")
code("""StrictMode > BrowserRouter > App > Routes
  AppLayout
    Sidebar > NavLink > Icon
    Header > Link > Icon
    main > Outlet
      DashboardPage > PageHeading StatCard Icon
      TasksPage > PageHeading Icon
      SettingsPage > PageHeading Icon
      NotFoundPage > PageHeading Link
    footer""", 9.2)
table(["Component", "Props", "Responsibility"], [
    ["App", "None", "Map route paths to page components"],
    ["AppLayout", "None", "Compose landmarks, navigation, Outlet and footer"],
    ["Sidebar", "None", "Brand, route links, active state and milestone context"],
    ["Header", "None", "Route breadcrumb and development setup link"],
    ["PageHeading", "eyebrow, title, children", "Consistent context, h1 and supporting copy"],
    ["StatCard", "label, value, detail, icon", "Reusable structural metric with descriptive context"],
    ["Icon", "name, size = 20", "Decorative local SVG; accessible meaning belongs to parent"],
    ["DashboardPage", "None", "Foundation overview, task preview and roadmap"],
    ["TasksPage", "None", "Describe the task feature boundary and later milestones"],
    ["SettingsPage", "None", "Commands, stack reference and setup disclosure"],
    ["NotFoundPage", "None", "Explain invalid routes and link back to Overview"],
    ["useRouteFocus", "Returns main ref", "Manage title and main focus after navigation"],
], [1.45, 1.75, 3.5], size=8.7)
h("Planned reusable task components")
p("TaskForm will accept initialValues, onSubmit, and onCancel. TaskList will receive tasks and callbacks for toggle, edit, and delete. TaskCard will render one task and its explicit controls. FilterBar will accept filters and onChange. A future EmptyState can take a title, description, and optional action. These contracts are plans; these components are not implemented in Week 1.")
p("The proposed task model contains id, title, description, status, priority, dueDate, createdAt, and updatedAt. Stable IDs will support lists and persistence. Filtered results should be derived from task state rather than stored as a second copy.")

# 6 — Build config snippets.
page("Webpack and Babel configuration")
p("webpack.config.cjs exports a function that selects development or production behavior from the CLI mode. These excerpts are from the actual configuration; the full file also defines HtmlWebpackPlugin, CSS extraction, splitChunks, runtimeChunk, and the local dev server.")
code("""const isProduction = argv.mode === 'production';
entry: './src/index.jsx',
output: {
  path: path.resolve(__dirname, 'dist'),
  filename: 'assets/[name].[contenthash:8].js',
  assetModuleFilename: 'assets/[name].[contenthash:8][ext]',
  publicPath: '/',
  clean: true,
},
resolve: { extensions: ['.js', '.jsx'] },
devtool: isProduction ? 'source-map' : 'eval-cheap-module-source-map',
module: {
  rules: [
    { test: /\\.jsx?$/, exclude: /node_modules/, use: 'babel-loader' },
    {
      test: /\\.css$/i,
      use: [
        isProduction ? MiniCssExtractPlugin.loader : 'style-loader',
        'css-loader',
      ],
    },
    { test: /\\.(svg|png|jpe?g|gif|webp|woff2?)$/i,
      type: 'asset/resource' },
  ],
},""", 8.4)
p("HtmlWebpackPlugin uses public/index.html and the local SVG favicon. The dev server binds to 127.0.0.1:3000, enables hot reload and historyApiFallback, and displays compile errors through the client overlay. Production builds emit hashed JavaScript, separate CSS, asset files, HTML, and source maps. Static hosting must provide an index.html fallback for application routes.")
h("Babel JSX compilation")
code((ROOT / "babel.config.cjs").read_text(), 8.7)
p("The automatic JSX runtime removes the need to import React in every JSX file. Environment transforms use Babel targets defaults. Browser support and polyfill needs should be revisited if future features require APIs outside the chosen browser baseline.")

# 7 — Quality config snippets.
page("ESLint and Prettier configuration")
p("The modern flat ESLint configuration scopes browser globals and JSX parsing to src files, uses Node globals for CommonJS configuration, and keeps formatting under Prettier. The actual file imports the packages shown here; its key rule composition is reproduced below.")
code("""import js from '@eslint/js';
import globals from 'globals';
import reactX from 'eslint-plugin-react-x';
import reactHooks from 'eslint-plugin-react-hooks';
import prettier from 'eslint-config-prettier';

// Inside the exported flat configuration array:
js.configs.recommended,
{
  files: ['src/**/*.{js,jsx}'],
  languageOptions: {
    globals: globals.browser,
    parserOptions: { ecmaFeatures: { jsx: true } },
  },
  plugins: { 'react-x': reactX, 'react-hooks': reactHooks },
  rules: {
    ...reactX.configs.recommended.rules,
    ...reactHooks.configs.recommended.rules,
    'no-unused-vars': ['error', { argsIgnorePattern: '^_' }],
    eqeqeq: ['error', 'always'],
  },
},
prettier,""", 8.5)
h("Shared formatting rules")
code((ROOT / ".prettierrc.json").read_text(), 9)
p("npm run lint executes eslint . --max-warnings 0. npm run format writes supported files, and npm run format:check checks them without changing content. Generated bundles, dependencies, evidence files, and DOCX/render artifacts are excluded from formatting. .editorconfig reinforces UTF-8, two spaces, LF endings, and a final newline.")
p("Configuration snippets are excerpts where indicated and must not replace the complete configuration files. Changes to lint plugins should be reviewed against their peer dependencies before updating the lockfile.")

# 8 — Standards and Git.
page("Coding standards and Git workflow")
h("Naming and organization")
table(["Area", "Convention", "Example"], [
    ["Components and JSX files", "PascalCase", "TaskCard.jsx"],
    ["Functions and variables", "camelCase", "formatDueDate"],
    ["Custom hooks", "use prefix and camelCase", "useRouteFocus.js"],
    ["Genuine constants", "UPPER_SNAKE_CASE", "STORAGE_VERSION"],
    ["CSS classes and docs", "lowercase with hyphens", "page-heading"],
    ["Module formats", ".jsx UI, .js helpers, .cjs build, .mjs ESLint", "webpack.config.cjs"],
], [1.7, 2.9, 2.1])
p("Use one primary component per file and extract shared UI only when its responsibility is clear. Keep state immutable and hooks unconditional. Use stable IDs for dynamic list keys. Prefer const, strict equality, derived values, and explicit callbacks. Describe prop contracts with JSDoc and architecture.md. Comments should explain a decision or constraint.")
h("Code quality practices")
p("Run lint, format check, and build before committing. Review the diff, keep package.json and package-lock.json synchronized, and avoid production debug logs. Do not commit credentials or assume frontend environment variables can keep secrets. The ignore rules exclude node_modules, dist, build, coverage, .env files, caches, and internal document render output while retaining genuine evidence logs.")
h("Git and GitHub")
p("A new repository was initialized inside the TaskFlow subdirectory, leaving the parent Expo project untouched. The local branch is codex/week-1. The implementation commit is c242254, titled feat: establish TaskFlow Week 1 React foundation. Documentation and evidence are committed separately at finalization.")
p("The GitHub connector authenticated as nitinmehra1017. The requested repository lookup returned 404, the connector had no repository creation tool, and browser access to GitHub was declined. No repository was created or pushed. The submission includes docs/github-setup.md with exact setup commands.")
code("""# After creating an empty PUBLIC repository on GitHub:
git remote add origin https://github.com/nitinmehra1017/taskflow-reactjs-internship.git
git push -u origin HEAD:main""", 8.1)
p("Future changes should use focused codex branches and pull requests. Explain the change and validation, review conflicts explicitly, and never force push shared work. If a remote already has commits, inspect it before integrating rather than replacing its history.")

# 9 — Testing and accessibility.
page("Testing and accessibility strategy")
h("Week 1 verification boundary")
p("Executed checks cover dependency installation, source quality, formatting, and development/production compilation. No Jest suite, React Testing Library suite, screen-reader audit, or browser interaction test was executed. A successful build proves compilation, not that the rendered interface is correct.")
h("Automated testing planned for Week 5")
p("Introduce Jest, babel-jest, jest-environment-jsdom, React Testing Library, user-event, and jest-dom. Reuse Babel's JSX configuration and add CSS and asset mocks. Prefer queries by role, label, and visible text rather than component internals or broad snapshots. Isolate browser storage and reset test state between cases.")
table(["Layer", "Planned checks"], [
    ["Pure helpers", "Title validation, filter/sort logic, date behavior and empty inputs"],
    ["Task components", "Accessible fields, validation errors, submit/cancel callbacks and disabled states"],
    ["Route integration", "Create, edit, complete and remove tasks; route navigation and recovery"],
    ["Persistence", "Hydration, reload retention, corrupt JSON, unavailable storage and quota failures"],
    ["Accessibility", "Names, roles, labels, focus movement and keyboard-only interactions"],
], [1.45, 5.25])
h("Accessibility built into the starter")
p("The shell uses semantic header, nav, main, aside, and footer landmarks, one h1 per route, native navigation links, a visible-on-focus skip link, visible focus outlines, and decorative SVGs hidden from assistive technology. NavLink marks active links. Route changes update document titles and move focus to main. CSS adapts at narrow widths and honors reduced-motion preferences.")
h("Manual acceptance strategy")
bullet("Navigate all routes using only Tab, Shift+Tab, Enter, and the skip link; verify focus order and visible indicators.")
bullet("Check a 390px viewport, 200% zoom, readable content, touch targets, and absence of horizontal overflow.")
bullet("Review contrast, screen-reader landmarks and link names, direct-route reload, wildcard recovery, and browser-console errors.")
p("These manual checks remain outstanding because browser access was declined. No full WCAG conformance claim is made. Future fields need persistent labels and connected errors; future dialogs need focus trapping, Escape handling, and focus restoration.")

# 10 — Roadmap.
page("Six week development roadmap")
p("Each week builds on the previous acceptance boundary. Week 1 establishes the environment and plan; later deliverables should remain incremental and reviewable. No Week 2 to Week 6 feature is represented as completed by this submission.")
table(["Week", "Focus", "Deliverables and acceptance"], [
    ["1", "Project setup and planning", "Custom build, reusable shell, routing, standards, architecture, report, evidence and Git commits"],
    ["2", "Task essentials", "Create, edit, complete and delete tasks; controlled form validation, stable IDs and empty/error states"],
    ["3", "State and storage", "Task provider/reducer, versioned local storage, hydration and storage failure recovery"],
    ["4", "Search and organization", "Search, status/priority filters, sorting, preferences and derived dashboard statistics"],
    ["5", "Testing and polish", "Jest and Testing Library setup, behavior tests, keyboard/contrast/zoom review and issue fixes"],
    ["6", "Final delivery", "Production hosting and SPA rewrites, performance review, browser acceptance, final docs and demonstration"],
], [0.6, 1.8, 4.3], size=9.7)
h("Delivery rules")
p("A feature is complete only when its intended user behavior works, its validation has actually run, and limitations are documented. Future dashboard values must be derived from real task state. Failure and empty states should be designed along with successful states. Keep dependencies and package-lock.json aligned at every milestone.")
h("Production considerations for Week 6")
p("The current build targets the origin root. A host must rewrite application routes to index.html. Subpath hosting needs coordinated changes to Webpack publicPath and BrowserRouter basename. Decide whether to publish source maps, evaluate cache headers for hashed assets, and review bundle size after features stabilize. Current bundle warnings remain visible rather than hidden by increasing thresholds.")
h("Change control")
p("If internship requirements change, update the roadmap and component contracts before broadening implementation. Authentication, backend APIs, and multi-user collaboration should not be added simply because directories can accommodate them. Keep the learning objectives and six week deadline visible in every weekly report.")

# 11 — Work breakdown.
page("Week 1 work breakdown")
p("The following allocation totals 35 hours and describes the intended Week 1 workload. It is a planning estimate, not a timesheet or a claim that this automated session lasted 35 hours. The intern should attach an actual attendance or activity log if YuvaIntern requires one.")
table(["Activity", "Hours", "Expected output"], [
    ["Requirements and scope analysis", "4", "Week boundaries, acceptance checklist and feature inventory"],
    ["Environment and dependency setup", "5", "Runtime preparation, package scripts, dependency installation and lockfile"],
    ["Webpack and Babel configuration", "6", "Development server, production bundle, HTML, CSS, assets and source maps"],
    ["Architecture and component planning", "5", "Folder ownership, prop contracts, hierarchy and future data flow"],
    ["Responsive shell and routing", "6", "Header, sidebar, Overview, Tasks preview, Settings and 404 route"],
    ["Standards and code quality tools", "3", "Flat ESLint, Prettier, naming conventions and review rules"],
    ["Verification and submission documentation", "6", "Executed checks, evidence notes, Word report and submission text"],
    ["Total planned allocation", "35", "Week 1 project setup and planning"],
], [2.5, 0.6, 3.6], size=10)
h("Work products available for review")
p("The source code, configuration, lockfile, architecture guide, standards guide, roadmap, publishing instructions, report, and submission text are concrete outputs. Evidence logs show what commands actually ran. The implementation and documentation are separated into meaningful local commits.")
h("Work that still needs access")
p("Browser acceptance checks and screenshots remain incomplete. A public GitHub repository must still be created and pushed using the supplied instructions or restored access. These gaps should be resolved before marking the entire internship submission checklist complete; they are not hidden by the 35 hour plan.")

# 12 — Setup and evidence.
page("Environment setup and verification evidence")
p("The workspace already contained an unrelated Expo project. TaskFlow was created as a separate subdirectory and local repository. The shell initially lacked node and npm on PATH. Bundled Node.js 24.19.0 was used, and npm 11.21.1 was installed in a temporary tools directory using bundled pnpm. TaskFlow itself uses npm and package-lock.json only.")
table(["Executed action", "Observed result", "Evidence file"], [
    ["npm install", "Succeeded; lint dependency migration recorded", "dependency-install.log"],
    ["npm ci with local cache", "Exit 0; 605 packages installed", "clean-install.log"],
    ["Runtime and npm ls", "Node 24.19.0, npm 11.21.1, Git 2.50.1; resolved tree", "environment.log"],
    ["npm run format", "Exit 0; supported files formatted", "format.log"],
    ["npm run lint", "Exit 0; zero warnings", "lint.log"],
    ["npm run format:check", "Exit 0; Prettier style confirmed", "format-check.log"],
    ["npm run build", "Exit 0; 3 bundle size/performance warnings", "build.log"],
    ["npm start", "Server announced port 3000 and compiled successfully", "dev-server.log"],
], [2.0, 2.9, 1.8], size=8.9)
h("Observed warnings and limits")
p("The initial sandbox server attempt failed with listen EPERM; an authorized retry started successfully. A sandbox HTTP probe could not reach the server. Browser-origin access to localhost was declined, so no successful HTTP response or browser rendering is claimed. The production build emitted a 250 KiB vendor asset and a 284 KiB entrypoint, triggering three standard Webpack size/performance warnings.")
p("A clean install reported a deprecated transitive uuid package and an optional fsevents install-script notice. Neither blocked compilation. No npm security audit was run. The final npm dependency tree resolved without invalid direct peers; broader vulnerability or multi-Node-version assurance is not claimed.")
h("Screenshots still required")
p("No screenshots are supplied. The local browser and GitHub origins were denied, and the native Terminal app was unavailable to computer use. Capture desktop/mobile Overview, Tasks/Settings routes, real setup output, lint/format/build output retaining warnings, and the project tree when access is available. evidence/README.md lists the exact capture and manual-check checklist.")

# 13 — Challenges and conclusion.
page("Challenges and conclusions")
h("Challenges and responses")
table(["Challenge", "Response and practical limit"], [
    ["Existing unrelated workspace", "Created an isolated TaskFlow subdirectory and Git repository; preserved all Expo files"],
    ["Missing shell runtime commands", "Used bundled Node and a temporary npm installation; documented normal Node 24 setup for the user"],
    ["Unsupported ESLint version", "Migrated to ESLint 10 and compatible React X/Hooks plugins, then completed a clean install and lint check"],
    ["Sandbox local port restriction", "Retried with required authorization; development compilation succeeded"],
    ["Browser and screenshot restrictions", "Did not circumvent denied access; retained real logs and marked UI verification and screenshots incomplete"],
    ["GitHub creation unavailable", "Prepared local commits and exact empty public repository setup/push instructions; no URL claimed"],
    ["Bundle size warnings", "Retained the warnings and scheduled performance review for Week 6"],
], [2.1, 4.6], size=9.3)
h("Conclusion")
p("Week 1 delivers a usable development foundation and a bounded plan for TaskFlow. The source compiles in development and production, lint and formatting checks pass, and the architecture provides clear places for future task behavior. Later weeks can now focus on user-facing task features instead of recreating setup decisions.")
p("The submission remains incomplete in three access-dependent areas: browser acceptance, genuine screenshots, and public GitHub publication. Complete those using the evidence checklist and publishing instructions. The 35 hour schedule is a planned allocation; it does not replace an actual attendance record.")
h("Primary references and companion files")
p("Technical references consulted on 19 October 2026: Node.js release support at nodejs.org/en/about/previous-releases; Webpack setup at webpack.js.org/guides/getting-started; ESLint flat configuration at eslint.org/docs/latest/use/configure/configuration-files. Installed package metadata and command logs establish the actual resolved versions.")
p("Companion files: README.md; docs/architecture.md; docs/coding-standards.md; docs/roadmap.md; docs/github-setup.md; evidence/README.md; submission-description.txt. The source of this report is docs/create_week1_report.py. The final DOCX is the Word file to submit.")

for element in [doc.element, doc.styles.element]:
    for border in element.xpath('.//w:pBdr'):
        border.getparent().remove(border)

doc.save(OUT)
print(f"Created {OUT}")
print(f"Paragraphs: {len(doc.paragraphs)}; tables: {len(doc.tables)}")
