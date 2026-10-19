"""Rebuild the Week 6 report using python-docx from the bundled document runtime."""

from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "ReactJS_Internship_Week6_Performance_and_Accessibility.docx"
doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.27)
section.page_height = Inches(11.69)
section.top_margin = Inches(0.78)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)
section.header_distance = Inches(0.3)
section.footer_distance = Inches(0.3)
section.different_first_page_header_footer = True

for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"]:
    style = doc.styles[name]
    style.font.name = "Arial"
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_after = Pt(8)

normal = doc.styles["Normal"]
normal.font.size = Pt(10.5)
normal.paragraph_format.line_spacing = 1.15
normal.paragraph_format.widow_control = True
for name, size in [("Title", 29), ("Heading 1", 21), ("Heading 2", 13), ("Heading 3", 11)]:
    doc.styles[name].font.size = Pt(size)
    doc.styles[name].font.bold = name != "Title"
    doc.styles[name].paragraph_format.keep_with_next = True
doc.styles["Subtitle"].font.size = Pt(13)
doc.styles["Subtitle"].font.italic = False

header = section.header.paragraphs[0]
header.text = "TaskFlow    /    YuvaIntern ReactJS Internship    /    Week 6"
header.runs[0].font.size = Pt(8)
header.runs[0].font.color.rgb = RGBColor(0, 0, 0)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run("Performance and accessibility    |    ").font.size = Pt(8)
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)

doc.core_properties.title = "TaskFlow ReactJS Internship Week 6 Performance and Accessibility"
doc.core_properties.subject = "YuvaIntern Junior ReactJS Developer Internship"
doc.core_properties.author = "TaskFlow project workspace"
doc.core_properties.keywords = "ReactJS, Webpack, Week 1, YuvaIntern, TaskFlow"


def p(text, style=None):
    return doc.add_paragraph(text, style=style)


def h(text, level=2):
    return doc.add_heading(text, level=level)


def page(title):
    doc.add_page_break()
    h(title, 1)


def bullet(text):
    par = p("•  " + text)
    par.paragraph_format.left_indent = Inches(0.1)
    par.paragraph_format.first_line_indent = Inches(-0.1)
    return par


def code(text, size=8.7):
    for line in text.strip("\n").splitlines():
        par = p(line or " ")
        par.paragraph_format.space_after = Pt(0)
        par.paragraph_format.line_spacing = 1.02
        par.paragraph_format.left_indent = Inches(0.1)
        par.paragraph_format.right_indent = Inches(0.08)
        par.paragraph_format.keep_with_next = False
        shade = OxmlElement("w:shd")
        shade.set(qn("w:fill"), "F2F4F5")
        par._p.get_or_add_pPr().append(shade)
        for run in par.runs:
            run.font.name = "Courier New"
            run.font.size = Pt(size)
    p("").paragraph_format.space_after = Pt(2)


def table(headers, rows, widths, size=9):
    tab = doc.add_table(rows=1, cols=len(headers))
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    tab.autofit = False
    for col, width in zip(tab.columns, widths):
        col.width = Inches(width)
    for cell, title, width in zip(tab.rows[0].cells, headers, widths):
        cell.width = Inches(width)
        cell.text = title
    repeat = OxmlElement("w:tblHeader")
    tab.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for cell, val, width in zip(tab.add_row().cells, row, widths):
            cell.width = Inches(width)
            cell.text = str(val)
    for ri, row in enumerate(tab.rows):
        cant_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant_split)
        for ci, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for side in ["top", "left", "bottom", "right"]:
                el = OxmlElement("w:" + side)
                el.set(qn("w:w"), "85" if side in ["top", "bottom"] else "95")
                el.set(qn("w:type"), "dxa")
                margins.append(el)
            props.append(margins)
            borders = OxmlElement("w:tcBorders")
            for side in ["top", "left", "bottom", "right"]:
                el = OxmlElement("w:" + side)
                el.set(qn("w:val"), "single")
                el.set(qn("w:sz"), "4")
                el.set(qn("w:color"), "D9D9D9")
                borders.append(el)
            props.append(borders)
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "243D48" if ri == 0 else ("F0F4F6" if ri % 2 == 0 else "FFFFFF"))
            props.append(shade)
            for par in cell.paragraphs:
                par.paragraph_format.space_after = Pt(0)
                par.paragraph_format.line_spacing = 1.08
                if widths[ci] < 0.8:
                    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in par.runs:
                    run.font.size = Pt(size)
                    run.font.bold = ri == 0
                    run.font.color.rgb = RGBColor(255, 255, 255) if ri == 0 else RGBColor(0, 0, 0)
    p("").paragraph_format.space_after = Pt(1)
    return tab



import json
C=json.loads((ROOT/'evidence/week6/comparison.json').read_text())
B=C['baseline']['bundle']; A=C['optimized']['bundle']
COV=json.loads((ROOT/'evidence/week6/coverage-summary.json').read_text())['total']
RESULT=json.loads((ROOT/'evidence/week6/jest-results.json').read_text())
EVIDENCE=ROOT/'evidence/week6'
def picture(filename,width,alt):
    par=p('');par.alignment=WD_ALIGN_PARAGRAPH.CENTER
    shape=par.add_run().add_picture(str(EVIDENCE/filename),width=Inches(width));shape._inline.docPr.set('descr',alt)
def caption(text):
    par=p(text)
    for run in par.runs:run.font.size=Pt(9);run.font.italic=True
p('YuvaIntern Junior ReactJS Developer Internship','Subtitle')
p('').paragraph_format.space_after=Pt(30)
p('TaskFlow Week 6 Performance and Accessibility','Title')
p('Final week implementation and verification report','Subtitle')
p('Week 6 of the six week ReactJS internship')
p('Prepared 19 October 2026')
h('Measured outcome')
p(f"Initial JavaScript decreased from {B['initialJsBytes']:,} to {A['initialJsBytes']:,} bytes, a {abs(C['changes']['initialJsPercent']):.2f} percent reduction. Initial CSS decreased from {B['initialCssBytes']:,} to {A['initialCssBytes']:,} bytes. A controlled React Profiler fixture reduced remote card update commits from eight to one. The full suite passed {RESULT['numPassedTests']} tests with zero failures.")
p('This report explains the baseline methodology, actual Lighthouse comparisons, route splitting, rendering improvements, responsive corrections and accessibility findings for internship review. Timing regressions and outstanding manual checks are included alongside improvements.')
p('Repository: github.com/nitinmehra1017/taskflow-reactjs-internship')
p('Final branch: week-6-performance-accessibility')
p('All previous reports and history are preserved. One entry-size build warning and 26 development dependency advisory entries remain. This report does not certify full WCAG compliance or hosted production readiness.')

page('Introduction objectives and measurement method')
p('TaskFlow combines a React dashboard, local task management and a Remote Task Explorer. Week 1 established custom Webpack, Babel, routing and architecture. Week 2 added seven reusable UI primitives and SCSS tokens. Week 3 added Context, reducers, custom hooks, CRUD, filters, statistics and localStorage recovery. Week 4 added validated public API requests and local imports. Week 5 added comprehensive tests, regression fixes and CI.')
p('Week 4 and Week 5 implementation was verified in the repository before branching from Week 5 commit 7af5394961c7cec85fdb1df4d01ec2423dd0e694. Week 6 retains the design system and custom Webpack architecture. TaskProvider remains above the persistent layout so route loading does not discard task state.')
h('Objectives')
bullet('Measure before changing performance source and repeat under the same lab method.')
bullet('Defer optional page code and avoid unnecessary remote card rendering.')
bullet('Fix actual contrast, naming, focus and reflow issues without removing features.')
bullet('Protect all previous workflows and provide traceable final documentation.')
h('Lighthouse method')
p('Twenty four baseline audits completed before performance source changes. The optimized phase repeats 24 audits: four routes, mobile and desktop presets, three runs each. Reports use Lighthouse 13.5.0, headless Chrome 155, default mobile simulated throttling and the desktop preset. Medians summarize three runs. Production assets are served at 127.0.0.1:4173 through the same gzip and cache-header preview server in both phases. Default Lighthouse storage/cache reset is retained.')
p('Local tasks start empty. Explore calls the real public API, whose response timing is uncontrolled. Local tool/CPU activity introduces noise. These observations are lab measurements, not field Core Web Vitals or proof of statistically significant speed gains. Raw JSON records settings, environment, audit details, requests and timestamps; HTML reports are retained.')

page('Webpack analysis and measured payload changes')
table(['Metric','Baseline','Optimized','Change'],[
('Initial JavaScript bytes',f"{B['initialJsBytes']:,}",f"{A['initialJsBytes']:,}",f"{C['changes']['initialJsPercent']}%"),
('Initial CSS bytes',f"{B['initialCssBytes']:,}",f"{A['initialCssBytes']:,}",f"{C['changes']['initialCssPercent']}%"),
('All emitted JavaScript bytes',f"{B['totalJsBytes']:,}",f"{A['totalJsBytes']:,}",'Includes deferred code'),
('Build warnings',B['warnings'],A['warnings'],'Entry-size remains')],[2.3,1.3,1.3,1.7],9)
p('These are uncompressed Webpack emitted bytes. Initial means the main entrypoint; a direct visit to a lazy route additionally downloads its page/shared chunks. Total emitted JavaScript can increase because code splitting needs runtime boundaries. Network transfer is measured separately in the Lighthouse tables.')
p('Production mode, tree shaking, source maps, runtime separation and splitChunks existed already. Week 6 adds named page chunks and meaningful React DOM/router cache groups. The oversized single vendor file is divided. The original 244 KiB warning budget is unchanged: the individual asset and generic recommendation warnings are gone, but one combined entrypoint warning remains.')
picture('payload-comparison.png',6.3,'Actual Webpack initial payload bytes and controlled React card update counts before and after optimization')
caption('Measured entry payload and Profiler fixture counts. Commit-count reduction is not a production CPU timing claim.')

page('Route loading rendering and API behavior')
p('Tasks, Components, Explore and Settings use module-scope React.lazy with named dynamic imports. Dashboard stays eager. AppLayout keeps navigation outside Suspense; only Outlet is suspended. A static minimum-height fallback provides a heading and polite status. It reserves space without animation, but variable page content can still shift; actual CLS is reported.')
code("const TasksPage = lazy(() =>\n  import(/* webpackChunkName: 'tasks' */ './pages/TasksPage')\n);\n<RouteErrorBoundary key={pathname} path={pathname}>\n  <Suspense fallback={<RouteLoading />}><Outlet /></Suspense>\n</RouteErrorBoundary>")
p('A rejected chunk reaches a route error boundary with a same-path full-document reload link. This supports stale-chunk recovery without automatic reload loops. The path key resets the boundary on navigation. Deferred/rejected module tests check fallback semantics, persistent navigation and recovery.')
h('React Profiler evidence')
p('Test-only instrumentation wraps the real Card in React Profiler, displays eight fixed remote records and imports record 1. Before the change all eight cards committed updates. After RemoteTaskCard memoization and a stable import callback, only record 1 committed: eight to one, an 87.5 percent reduction for this fixture. A regression test now requires that single update. Raw development/jsdom durations are retained without claiming production speedup.')
p('Existing TaskCard memoization, stable task callbacks, memoized statistics, provider value and local filters were retained. No blanket memoization or unnecessary context rewrite was added.')
h('API review')
p('Existing parallel todos/users requests, response validation, eight-record pagination, AbortController cleanup, timeout, stale-response guards and retained results during refresh were retested. Search/filter/page changes do not refetch: the existing test checks two endpoint calls. Revisiting Explore intentionally fetches fresh data. Week 6 normalizes search once per calculation and skips unchanged remote-card renders. No cache changes freshness semantics.')

for mode in ['mobile','desktop']:
    page(mode.title()+' Lighthouse median comparison')
    rows=[]
    for b,a in zip(C['baseline']['lighthouse'],C['optimized']['lighthouse']):
        if b['mode']!=mode:continue
        rows.append((b['route'].title(),f"{b['performance']:.0f} to {a['performance']:.0f}",f"{b['accessibility']:.0f} to {a['accessibility']:.0f}",f"{b['fcpMs']:.0f} to {a['fcpMs']:.0f}",f"{b['lcpMs']:.0f} to {a['lcpMs']:.0f}",f"{b['speedIndexMs']:.0f} to {a['speedIndexMs']:.0f}"))
    p('Three runs per route. Scores use a 0 to 100 scale. Timings are milliseconds; each cell shows baseline to optimized. Raw reports are in the matching baseline and optimized evidence folders.')
    table(['Route','Perf','A11y','FCP ms','LCP ms','Speed Index'],rows,[1.05,.8,.8,1.25,1.25,1.45],8.5)
    rows=[]
    for b,a in zip(C['baseline']['lighthouse'],C['optimized']['lighthouse']):
        if b['mode']!=mode:continue
        rows.append((b['route'].title(),f"{b['tbtMs']:.0f} to {a['tbtMs']:.0f}",f"{b['cls']:.4f} to {a['cls']:.4f}",f"{b['transferredBytes']:,.0f} to {a['transferredBytes']:,.0f}",f"{b['requests']:.0f} to {a['requests']:.0f}"))
    table(['Route','TBT ms','CLS','Transfer bytes','Requests'],rows,[1.05,1,1.3,2.25,1],8.5)
    h('How to interpret these results')
    p('Scores and timings can improve, regress or remain unchanged. Lower payload does not prove lower LCP in every run. Network latency, CPU variation, browser scheduling and additional route requests affect measurements. Explore has uncontrolled public API timing. No feature was removed or audit disabled to improve the scores.')
    p('The full numeric comparison and percentage changes are in evidence/week6/comparison.json and docs/Week6_Performance_and_Accessibility.md. A zero baseline has no percent-change value; it is not reported as an infinite improvement. Request counts include the resources Lighthouse observed during initial load.')

page('Accessibility findings and WCAG evaluation')
p('Real-browser Lighthouse axe checks supplement Jest/jsdom semantics and manual keyboard observations. No full WCAG 2.2 AA compliance claim is made. Four genuine issue groups were corrected; findings and verification are retained.')
table(['Issue and criterion','Finding and change','Verification'],[
('Text contrast 1.4.3','Muted breadcrumb, eyebrow, footer, welcome, roadmap and remote intro text failed contrast. Darkened within the existing palette.','Baseline selector findings, calculated ratios and final browser audits'),
('Reflow 1.4.10','Settings width 365 at viewport 320. Grid minimum and nonwrapping badge fixed; command text wraps.','Final width 320; 35 route/width checks'),
('Label in Name 2.5.3','Import local copy text was missing from its accessible name. Visible text now matches Import remote task plus ID.','Dedicated regression and browser label audit'),
('Focus 1.4.11 and 2.4.7','Green outline had 2.41 to 1 contrast on forest. Contextual lime ring now 9.11 to 1.','Computed 3px lime outline and keyboard check')],[1.55,3.15,1.9],8.7)
h('Additional criteria evaluated')
p('Semantic structure, landmarks and control names (1.3.1, 4.1.2); written non-color task statuses (1.4.1); actual text zoom (1.4.4); keyboard operation and traps (2.1.1, 2.1.2); skip link (2.4.1); focus order and visibility (2.4.3, 2.4.7); dialog bounds and obscuration (2.4.11); target sizes (2.5.8); form identification/labels (3.3.1, 3.3.2); and status messages (4.1.3) were reviewed through the documented automated and browser checks.')
p('Screen-reader/speech-recognition software, independent engines, exhaustive interaction-state contrast and reduced-motion preference emulation remain unverified. Existing reduced-motion CSS and static fallback were inspected. Native select arrow selection did not change value through automation; it is recorded as a manual verification gap, while automated filter selection tests pass.')

page('Responsive layouts and keyboard verification')
p('All five major routes were checked at 320, 375, 390, 768, 1024, 1440 and 1920 CSS pixels: 35 final measurements without horizontal overflow. Every route heading was awaited. The 320px Settings overflow was fixed after an intermediate 339px result exposed a remaining nonwrapping badge.')
p('Native Chrome confirmed 200 percent zoom; all five routes fit the 735 CSS pixel effective viewport. Zoom was restored to 100 percent. Below 400px the dashboard now uses two statistic columns; the sidebar can scroll in short desktop viewports. A 320 by 700 dialog measured left/top 12 and right/bottom 308/688 with internal content scrolling.')
p('Keyboard checks created a task, triggered invalid form feedback, edited its title, used Space to complete and Enter to reopen, searched and cleared, then confirmed deletion with Tab and Return. Only the newly created QA task was deleted. Shift Tab wrapped Close to Cancel. Escape restored dialog triggers. Explorer search, clear and pagination focused the results heading. The detailed log records the native select limitation.')
picture('optimized/dialog-320.jpg',2.25,'Actual task dialog at 320 by 700 CSS pixels with internal scrolling')
caption('Actual 320px task dialog. Its boundaries remain inside the viewport; fields and actions can be reached by internal scrolling.')

page('Genuine application and audit captures')
picture('optimized/explore-final-1440.jpg',5.6,'Actual optimized production remote explorer with matching import button labels')
caption('Final production Explorer showing live records and the corrected visible import labels.')
picture('optimized/lighthouse-dashboard-mobile.jpg',5.6,'Actual saved Lighthouse dashboard mobile run displayed in Chrome')
caption('One genuine Lighthouse run. The report tables summarize three-run medians rather than treating this single capture as the median.')

page('Regression quality security and challenges')
p(f"The final coverage execution passed {RESULT['numPassedTests']} tests across {RESULT['numPassedTestSuites']} suites with zero failures. All 111 previous assertions remain. New tests cover lazy resolution, pending/error semantics, recovery links, split-route navigation and retained state, import label naming and the Profiler update count.")
table(['Coverage metric','Actual','Threshold'],[(key.title(),str(COV[key]['pct'])+'%','80%') for key in ['statements','branches','functions','lines']],[3.2,1.7,1.7])
p('Existing tests now await actual lazy modules through asynchronous act; production route modules are not mocked away. A deferred API response preserves the original loading-state assertion. CSS layout claims come from browser measurements rather than jsdom. npm ci, Jest, ESLint, formatting and production build passed. The production build retains one combined entry-size warning. GitHub Actions runs installation, lint, formatting, coverage and build on week branches.')
h('Dependency review')
p('Before and after a compatible npm audit fix, the full graph reports 26 development affected-package entries: 21 moderate and five high, zero critical. Production-only audit reports zero vulnerabilities. No compatible update cleared the remaining Webpack development server and Jest transitive chains. No force fix was used. The lockfile change records project version 0.6.0. These findings remain maintenance work before deployment tooling approval.')
h('Challenges and solutions')
p('Route splitting required asynchronous test setup without weakening behavior assertions. Browser audits found contrast and label-in-name issues missed by jsdom. A first reflow fix did not address the Settings badge, so actual widths were rechecked and corrected. Final audit runs were restarted after source fixes. Timing noise and public API latency require cautious comparison; regressions are retained.')
p('Independent browser engines, assistive-technology checks, reduced-motion emulation, native-select arrow verification, one build warning and development advisories remain. The loopback preview is a verification server, not a hosted deployment.')

page('Planned effort final readiness and conclusion')
p('The following allocation totals 35 planned hours. It is a work breakdown, not a measured attendance record.')
table(['Activity','Planned hours'],[
('Verify previous work and prepare branch',2),('Baseline and repeated performance measurement',4),('Webpack route splitting and loading recovery',5),('Profiler analysis and rendering optimization',4),('Responsive layout and zoom verification',5),('Accessibility evaluation and genuine fixes',5),('Automated regression testing',3),('Security review and CI validation',2),('Documentation evidence and final reports',5),('Total',35)],[5.3,1.3])
h('Final internship readiness')
p('All six weeks of source work and Word reports are present in this branch. Week 1 setup, Week 2 component library, Week 3 task manager, Week 4 remote integration and Week 5 QA are retained. Week 6 adds measured payload reduction, targeted rendering work, responsive and accessibility fixes and the final summary. docs/Final_Internship_Project_Summary.md links the six weeks and explains architecture, commands and limitations.')
p('The project is ready for internship review with the documented gaps. Main has not been merged and no deployment is claimed. Future work should address independent engines, assistive technologies, field monitoring, bundle budgets and compatible dependency migrations.')
h('Install run and submit')
p('Use Node 24 and npm 10 or later. Run npm ci, then npm start for localhost:3000. For production verification run npm run build and npm run preview, then open localhost:4173. Submit this Word report, the week-6-performance-accessibility branch URL and submission-week6-description.txt.')
for element in list(doc.element.iter(qn('w:pBdr'))):element.getparent().remove(element)
for style in doc.styles:
    for element in list(style.element.iter(qn('w:pBdr'))):element.getparent().remove(element)
doc.core_properties.keywords='ReactJS, YuvaIntern, TaskFlow, Week 6, performance, accessibility'
doc.save(OUT)
print(OUT)
